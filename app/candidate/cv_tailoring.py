from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from app.candidate.resume_parser import extract_resume_text, find_latest_resume


DATA_ENGINEERING_CV = "DATA_ENGINEERING_CV"
AI_DATA_ENGINEERING_CV = "AI_DATA_ENGINEERING_CV"
BI_DATA_ANALYTICS_CV = "BI_DATA_ANALYTICS_CV"

SKILL_GROUPS = {
    "BI & Reporting": ["Oracle BI (OBIEE)", "ODI", "Power BI", "Tableau"],
    "Data Engineering": [
        "Oracle Data Integrator (ODI)",
        "ETL Pipelines",
        "Data Modeling (Physical, BMM, Presentation Layers)",
        "Data Validation",
        "Query Optimization",
    ],
    "Databases": ["Oracle DB", "MySQL", "PostgreSQL", "Qdrant"],
    "Programming": ["SQL", "PL/SQL", "Python (Pandas, NumPy)", "C++"],
    "Platforms": ["Linux", "Windows", "LangChain"],
    "Tools": ["Docker", "Git/GitHub", "Jupyter Notebook", "VS Code"],
    "AI Engineering": [
        "Retrieval-Augmented Generation (RAG)",
        "Vector Databases",
        "Semantic Search",
        "Embeddings",
        "Prompt Engineering",
    ],
}

EXPERIENCE = [
    {
        "title": "AI Engineer",
        "period": "June 2026 - Present",
        "employer": "X VIEW (CESI SpA)",
        "location": "Milan, Italy",
        "bullets": [
            "Design and develop Python-based document ingestion pipelines for an enterprise semantic search and knowledge retrieval platform.",
            "Develop a REST API connecting Oracle APEX with AI-powered semantic search services, enabling integration between enterprise applications and data services.",
            "Implement document processing workflows covering extraction, chunking, embedding generation, metadata enrichment, and vector indexing.",
            "Integrate Qdrant Vector Database to support semantic similarity search and scalable retrieval of enterprise knowledge.",
            "Integrate Azure OpenAI embedding services for enterprise documents and natural-language queries.",
        ],
    },
    {
        "title": "Data Analyst / BI Consultant",
        "period": "January 2026 - Present",
        "employer": "X VIEW (Poste Italiane)",
        "location": "Milan, Italy",
        "bullets": [
            "Design and maintenance of ETL/ELT processes using Oracle Data Integrator (ODI), including Mappings, Packages, Variables, Models, and Knowledge Modules.",
            "Develop data transformation and aggregation processes within Oracle Database to support enterprise Data Warehouse and analytical workloads.",
            "Perform SQL and PL/SQL development, debugging, and query optimization using Oracle SQL Developer.",
            "Analyze and validate DWH reconciliation and quadrature processes, improving data consistency, reliability, and trust in business-critical datasets.",
            "Identify and investigate data anomalies and inconsistencies across ETL and reporting workflows through SQL-based root-cause analysis.",
        ],
    },
    {
        "title": "Business Data Analyst / Oracle BI Specialist",
        "period": "January 2025 - January 2026",
        "employer": "X VIEW (Telecom Italia Sparkle)",
        "location": "Milan, Italy",
        "bullets": [
            "Migrated 150+ enterprise telecom reports from Hyperion to Oracle BI (OBIEE) for TIM Sparkle.",
            "Designed and optimized data models across Physical, BMM, and Presentation layers to support business reporting.",
            "Improved ETL workflows and schema performance, reducing refresh time and enhancing data accessibility.",
            "Scheduling and monitoring reports using Oracle BI Job Manager.",
        ],
    },
]

PROJECTS = [
    {
        "title": "StoreSight - Sales Forecasting System",
        "technologies": "Python, Flask, Docker, Machine Learning",
        "bullets": [
            "Developed a machine learning solution for retail sales forecasting using historical sales data.",
            "Performed data preprocessing and feature engineering to prepare datasets for predictive modelling.",
            "Implemented time-series forecasting techniques to support prediction of future sales trends.",
            "Developed a Flask-based backend and containerized the application using Docker.",
        ],
    },
    {
        "title": "Airbnb Amsterdam - Data Engineering & Analytics",
        "technologies": "Python, ETL, MySQL, Tableau",
        "bullets": [
            "Developed Python-based ETL pipelines to transform and load heterogeneous datasets into MySQL.",
            "Performed data preparation and transformation to support downstream analytical workloads.",
            "Created Tableau dashboards to analyze pricing, availability, and geographic patterns.",
        ],
    },
]

EDUCATION = [
    {
        "institution": "Sapienza University of Rome",
        "location": "Rome, Italy",
        "degree": "Master of Science in Engineering in Computer Science",
        "period": "2021 - 2025",
        "details": [
            "Thesis: Quantitative Analysis of Secure Code Generation by LLMs and Generative AI, focusing on vulnerability detection and mitigation in AI-generated software."
        ],
    },
    {
        "institution": "University of Mumbai",
        "location": "Mumbai, India",
        "degree": "Bachelor's Degree in Computer Science",
        "period": "2017 - 2020",
        "details": [],
    },
]

CERTIFICATIONS = [
    "IELTS English - C1 Proficient, British Council",
    "IBM Data Science Specialization, Coursera",
    "Machine Learning with Python, Coursera",
    "Databases and SQL for Data Science, Coursera",
    "Data Visualization with Python, Coursera",
    "Applied Data Science Capstone, Coursera",
]

PROFILES = {
    DATA_ENGINEERING_CV: {
        "headline": "DATA ENGINEERING | ETL | DATA WAREHOUSING",
        "summary": (
            "Data and BI professional with hands-on experience designing Oracle Data Integrator ETL/ELT processes, "
            "developing SQL and PL/SQL, transforming data for enterprise data warehouses, and validating DWH reconciliation. "
            "Additional Python ETL and analytics experience through project work."
        ),
        "skill_order": ["Data Engineering", "Programming", "Databases", "BI & Reporting", "Tools", "Platforms", "AI Engineering"],
        "project_order": [1, 0],
    },
    AI_DATA_ENGINEERING_CV: {
        "headline": "AI ENGINEERING | SEMANTIC SEARCH | DATA PIPELINES",
        "summary": (
            "AI Engineer with experience building Python document-ingestion and semantic-search workflows, integrating REST APIs, "
            "Qdrant vector search, and Azure OpenAI embeddings. Professional experience also includes Oracle ODI, data warehousing, "
            "SQL/PLSQL, and enterprise BI."
        ),
        "skill_order": ["AI Engineering", "Programming", "Databases", "Data Engineering", "Platforms", "Tools", "BI & Reporting"],
        "project_order": [0, 1],
    },
    BI_DATA_ANALYTICS_CV: {
        "headline": "BUSINESS DATA ANALYSIS | ORACLE BI | REPORTING",
        "summary": (
            "Business Data Analyst and Oracle BI specialist with experience migrating 150+ enterprise reports from Hyperion to OBIEE, "
            "designing Physical, BMM, and Presentation-layer data models, optimizing ETL workflows, and supporting enterprise reporting. "
            "Also experienced in SQL/PLSQL development and data validation."
        ),
        "skill_order": ["BI & Reporting", "Data Engineering", "Programming", "Databases", "Tools", "Platforms", "AI Engineering"],
        "project_order": [1, 0],
    },
}


def select_cv_version(job: dict[str, Any] | str) -> str:
    if isinstance(job, str):
        title = job
        technologies: list[str] = []
    else:
        title = str(job.get("title") or "")
        technologies = [str(skill) for skill in job.get("technologies", [])]

    title_text = title.lower()
    skill_text = " ".join(technologies).lower()
    if any(term in title_text for term in ("ai engineer", "machine learning", "forward deployed ai")):
        return AI_DATA_ENGINEERING_CV
    if any(term in skill_text for term in ("rag", "qdrant", "semantic search", "vector database", "azure openai", "embeddings")):
        return AI_DATA_ENGINEERING_CV
    if any(term in title_text for term in ("bi", "business intelligence", "oracle bi", "reporting analyst")):
        return BI_DATA_ANALYTICS_CV
    if "data analyst" in title_text:
        return BI_DATA_ANALYTICS_CV
    return DATA_ENGINEERING_CV


def _normalize_for_validation(text: str) -> str:
    return re.sub(r"\s+", " ", text.lower()).strip().rstrip(" .")


def build_tailored_cv(resume_path: str | Path, version: str) -> dict[str, Any]:
    if version not in PROFILES:
        raise ValueError(f"Unknown CV profile: {version}")

    resume_file = find_latest_resume(resume_path)
    source = extract_resume_text(resume_file)
    lines = [line.strip() for line in source.splitlines() if line.strip()]
    if not lines:
        raise ValueError("The source CV contains no extractable text.")

    source_normalized = _normalize_for_validation(source)
    profile = PROFILES[version]
    experiences = [dict(role) for role in EXPERIENCE]
    projects = [dict(PROJECTS[index]) for index in profile["project_order"]]
    for role in experiences:
        role["bullets"] = list(role["bullets"])
    for project in projects:
        project["bullets"] = list(project["bullets"])

    factual_bullets = [
        bullet
        for role in experiences
        for bullet in role["bullets"]
    ] + [
        bullet
        for project in projects
        for bullet in project["bullets"]
    ]
    missing_from_source = [
        bullet for bullet in factual_bullets
        if _normalize_for_validation(bullet) not in source_normalized
    ]
    if missing_from_source:
        raise ValueError("Tailored CV content does not match source CV text: " + missing_from_source[0])

    return {
        "version": version,
        "source_path": str(resume_file),
        "name": lines[0],
        "contact": lines[1] if len(lines) > 1 else "",
        "headline": profile["headline"],
        "summary": profile["summary"],
        "skills": [
            {"category": category, "items": list(SKILL_GROUPS[category])}
            for category in profile["skill_order"]
        ],
        "experience": experiences,
        "projects": projects,
        "education": [dict(item) for item in EDUCATION],
        "certifications": list(CERTIFICATIONS),
    }


def write_tailored_cv(resume_path: str | Path, version: str, output_path: str | Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    cv = build_tailored_cv(resume_path, version)
    target = Path(output_path)
    target.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor(38, 48, 57)
    normal.paragraph_format.space_after = Pt(2)

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(2)
    run = name.add_run(cv["name"])
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(23)
    run.font.color.rgb = RGBColor(22, 58, 79)

    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(3)
    contact_run = contact.add_run(cv["contact"])
    contact_run.font.size = Pt(8.5)
    contact_run.font.color.rgb = RGBColor(83, 96, 104)

    headline = document.add_paragraph()
    headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    headline.paragraph_format.space_after = Pt(5)
    headline_run = headline.add_run(cv["headline"])
    headline_run.bold = True
    headline_run.font.size = Pt(10)
    headline_run.font.color.rgb = RGBColor(13, 112, 115)

    summary = document.add_paragraph(cv["summary"])
    summary.paragraph_format.space_after = Pt(5)

    def add_section(title: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.space_after = Pt(2)
        heading = paragraph.add_run(title.upper())
        heading.bold = True
        heading.font.size = Pt(10.5)
        heading.font.color.rgb = RGBColor(13, 112, 115)

    def add_bullet(text: str) -> None:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Inches(0.18)
        paragraph.paragraph_format.first_line_indent = Inches(-0.12)
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.add_run(text)

    add_section("Technical Skills")
    for group in cv["skills"]:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        label = paragraph.add_run(group["category"] + ": ")
        label.bold = True
        paragraph.add_run(", ".join(group["items"]))

    add_section("Professional Experience")
    for role in cv["experience"]:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(3)
        heading.paragraph_format.space_after = Pt(1)
        title_run = heading.add_run(role["title"])
        title_run.bold = True
        heading.add_run(f" | {role['employer']} | {role['location']} | {role['period']}")
        for bullet in role["bullets"]:
            add_bullet(bullet)

    add_section("Selected Projects")
    for project in cv["projects"]:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(2)
        heading.paragraph_format.space_after = Pt(1)
        title_run = heading.add_run(project["title"])
        title_run.bold = True
        heading.add_run(" | " + project["technologies"])
        for bullet in project["bullets"]:
            add_bullet(bullet)

    add_section("Education")
    for education in cv["education"]:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        title_run = paragraph.add_run(education["institution"])
        title_run.bold = True
        paragraph.add_run(f" | {education['location']} | {education['period']}")
        detail = document.add_paragraph(education["degree"])
        detail.paragraph_format.space_after = Pt(1)
        for item in education["details"]:
            add_bullet(item)

    add_section("Certifications")
    for certification in cv["certifications"]:
        add_bullet(certification)

    document.save(target)
    return target


def generate_all_tailored_cvs(resume_path: str | Path, output_dir: str | Path) -> list[Path]:
    output_dir = Path(output_dir)
    return [
        write_tailored_cv(resume_path, version, output_dir / f"{version}.docx")
        for version in PROFILES
    ]
